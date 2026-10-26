import inspect
import textwrap
import unittest
from collections import defaultdict
from pathlib import Path

from django.apps import apps
from django.conf import settings
from django.core.management.base import BaseCommand
from docx import Document
from docx.shared import Pt


def iter_tests(suite):
    """Рекурсивно достаёт отдельные тесты из TestSuite."""
    for item in suite:
        if isinstance(item, unittest.TestSuite):
            yield from iter_tests(item)
        else:
            yield item


class Command(BaseCommand):
    help = "Записывает все тесты проекта в Word-файл"

    def add_arguments(self, parser):
        parser.add_argument(
            "-o", "--output",
            default=str(Path(settings.BASE_DIR) / "tests_report.docx"),
            help="Путь к итоговому .docx файлу",
        )

    def handle(self, *args, **options):
        base_dir = Path(settings.BASE_DIR).resolve()
        loader = unittest.TestLoader()

        # app -> класс -> список методов
        collected = defaultdict(lambda: defaultdict(list))

        for app in apps.get_app_configs():
            app_path = Path(app.path).resolve()
            # берём только свои приложения, а не django.contrib.*
            if base_dir not in app_path.parents and app_path != base_dir:
                continue

            suite = loader.discover(
                start_dir=str(app_path),
                pattern="test*.py",
                top_level_dir=str(base_dir),
            )
            for test in iter_tests(suite):
                cls = test.__class__
                method_name = getattr(test, "_testMethodName", None)
                try:
                    method = getattr(cls, method_name)
                except (AttributeError, TypeError):
                    continue  # тест, который не удалось импортировать
                collected[app.label][cls].append((method_name, method))

        doc = Document()
        doc.add_heading("Отчёт по тестам проекта", level=0)

        total = 0
        for app_label, classes in collected.items():
            doc.add_heading(f"Приложение: {app_label}", level=1)

            for cls, methods in classes.items():
                doc.add_heading(cls.__name__, level=2)
                if cls.__doc__:
                    doc.add_paragraph(inspect.cleandoc(cls.__doc__))

                for name, method in sorted(methods):
                    total += 1
                    doc.add_heading(name, level=3)

                    if method.__doc__:
                        doc.add_paragraph(inspect.cleandoc(method.__doc__))

                    try:
                        source = textwrap.dedent(inspect.getsource(method))
                    except (OSError, TypeError):
                        source = "(исходный код недоступен)"

                    p = doc.add_paragraph()
                    run = p.add_run(source)
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)

        doc.add_paragraph(f"Всего тестов: {total}")

        output = Path(options["output"])
        doc.save(output)
        self.stdout.write(self.style.SUCCESS(f"Готово: {output} (тестов: {total})"))